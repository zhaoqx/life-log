# -*- coding: utf-8 -*-
"""
简历导出模块 - Word和PDF
"""
import os
import sys
from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_background(cell, color):
    """设置单元格背景色"""
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_run_font(run, font_name='微软雅黑', font_size=11, bold=False, color=None):
    """设置字体"""
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(font_size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(color[0], color[1], color[2])

def add_section_title(doc, title, color=(139, 69, 19)):
    """添加章节标题（棕色系，模仿模板1）"""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(16)
    para.paragraph_format.space_after = Pt(8)
    run = para.add_run(f'◆ {title}')
    set_run_font(run, font_name='黑体', font_size=14, bold=True, color=color)
    return para

def export_to_word(student, base_dir):
    """导出为Word文档"""
    doc = Document()
    
    # 设置页面
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
    
    # ========== 封面 ==========
    # 标题
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(80)
    run = para.add_run('个 人 简 历')
    set_run_font(run, font_name='华文行楷', font_size=48, bold=True, color=(139, 69, 19))
    
    # 副标题
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(20)
    run = para.add_run('PERSONAL RESUME')
    set_run_font(run, font_size=14, color=(150, 150, 150))
    
    # 装饰线
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(40)
    run = para.add_run('─' * 35)
    set_run_font(run, font_size=10, color=(200, 180, 160))
    
    # 封面信息
    cover_table = doc.add_table(rows=4, cols=2)
    cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover_info = [
        ('姓    名', student.name or ''),
        ('所在学校', student.school or ''),
        ('联系电话', student.phone or '（请填写）'),
        ('投报学校', '（请填写）'),
    ]
    for i, (label, value) in enumerate(cover_info):
        row = cover_table.rows[i]
        cell0 = row.cells[0]
        cell0.text = ''
        para0 = cell0.paragraphs[0]
        para0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run0 = para0.add_run(f'{label}：')
        set_run_font(run0, font_size=14, bold=True, color=(139, 69, 19))
        
        cell1 = row.cells[1]
        cell1.text = ''
        para1 = cell1.paragraphs[0]
        run1 = para1.add_run(value)
        set_run_font(run1, font_size=14)
    
    # 日期
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(80)
    run = para.add_run(f'{datetime.now().year}年{datetime.now().month}月')
    set_run_font(run, font_size=12, color=(100, 100, 100))
    
    doc.add_page_break()
    
    # ========== 基本信息页 ==========
    add_section_title(doc, '基本信息')
    
    # 基本信息表格（带照片）
    info_table = doc.add_table(rows=5, cols=5)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # 设置列宽
    widths = [Cm(2.5), Cm(4), Cm(2.5), Cm(3.5), Cm(3.5)]
    for i, width in enumerate(widths):
        for cell in info_table.columns[i].cells:
            cell.width = width
    
    info_data = [
        ('姓    名', student.name or '', '性    别', student.gender or ''),
        ('出生日期', student.birth_date or '', '身    高', student.height or ''),
        ('所在小学', student.school or '', '班    级', student.class_name or ''),
        ('户籍所在', student.district or '', '民    族', '汉族'),
        ('联系电话', student.phone or '', '家庭住址', student.address or ''),
    ]
    
    for i, (l1, v1, l2, v2) in enumerate(info_data):
        row = info_table.rows[i]
        for j, (label, value) in enumerate([(l1, v1), (l2, v2)]):
            idx = j * 2
            # 标签
            cell = row.cells[idx]
            cell.text = ''
            para = cell.paragraphs[0]
            run = para.add_run(label)
            set_run_font(run, font_size=11, bold=True, color=(139, 69, 19))
            set_cell_background(cell, 'FFF5EE')
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            
            # 值
            cell = row.cells[idx + 1]
            cell.text = ''
            para = cell.paragraphs[0]
            run = para.add_run(value)
            set_run_font(run, font_size=11)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    
    # 合并照片单元格
    photo_cell = info_table.cell(0, 4)
    for i in range(1, 5):
        photo_cell.merge(info_table.cell(i, 4))
    
    # 添加照片
    photo_para = photo_cell.paragraphs[0]
    photo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if student.photo:
        photo_path = base_dir / 'static' / student.photo
        if photo_path.exists():
            run = photo_para.add_run()
            run.add_picture(str(photo_path), width=Cm(3))
    else:
        run = photo_para.add_run('[ 照片 ]')
        set_run_font(run, font_size=12, color=(150, 150, 150))
    set_cell_background(photo_cell, 'FFF8F0')
    
    # ========== 兴趣爱好 ==========
    if student.hobbies:
        add_section_title(doc, '兴趣爱好')
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Cm(0.5)
        run = para.add_run(student.hobbies)
        set_run_font(run, font_size=11)
    
    # ========== 荣誉奖励 ==========
    if student.awards:
        add_section_title(doc, '荣誉奖励')
        
        # 按类别分组
        categories = {}
        for award in student.awards:
            cat = award.category or '其他'
            if cat not in categories:
                categories[cat] = []
            categories[cat].append(award)
        
        for cat, awards in categories.items():
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.3)
            para.paragraph_format.space_before = Pt(6)
            run = para.add_run(f'【{cat}】')
            set_run_font(run, font_size=11, bold=True, color=(139, 90, 43))
            
            for award in sorted(awards, key=lambda x: x.order):
                para = doc.add_paragraph()
                para.paragraph_format.left_indent = Cm(0.5)
                para.paragraph_format.space_before = Pt(2)
                
                if award.date:
                    run = para.add_run(f'◆ {award.date}  ')
                    set_run_font(run, font_size=10, color=(100, 100, 100))
                
                run = para.add_run(award.title)
                set_run_font(run, font_size=11)
                
                if award.description:
                    para2 = doc.add_paragraph()
                    para2.paragraph_format.left_indent = Cm(1)
                    run2 = para2.add_run(award.description)
                    set_run_font(run2, font_size=10, color=(100, 100, 100))
    
    # ========== 自我介绍 ==========
    if student.self_intro:
        add_section_title(doc, '自我介绍')
        para = doc.add_paragraph()
        run = para.add_run(student.self_intro)
        set_run_font(run, font_size=11)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    
    # ========== 家长寄语 ==========
    if student.parent_message:
        add_section_title(doc, '家长寄语')
        para = doc.add_paragraph()
        run = para.add_run(student.parent_message)
        set_run_font(run, font_size=11)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    
    # ========== 老师评语 ==========
    if student.teacher_comment:
        add_section_title(doc, '老师评语')
        para = doc.add_paragraph()
        run = para.add_run(student.teacher_comment)
        set_run_font(run, font_size=11)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    
    # ========== 证书照片页 ==========
    if student.certificates:
        doc.add_page_break()
        add_section_title(doc, '证书与荣誉')
        
        # 每行放2张证书
        certs = list(student.certificates)
        for i in range(0, len(certs), 2):
            table = doc.add_table(rows=2, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            for j in range(2):
                if i + j < len(certs):
                    cert = certs[i + j]
                    # 图片单元格
                    cell = table.cell(0, j)
                    para = cell.paragraphs[0]
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    
                    if cert.image_path:
                        img_path = base_dir / 'static' / cert.image_path
                        if img_path.exists():
                            run = para.add_run()
                            run.add_picture(str(img_path), width=Cm(7))
                    
                    # 标题单元格
                    cell = table.cell(1, j)
                    para = cell.paragraphs[0]
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = para.add_run(cert.name)
                    set_run_font(run, font_size=10, color=(100, 100, 100))
            
            # 添加间距
            doc.add_paragraph()
    
    # 保存
    output_dir = base_dir / 'static' / 'exports'
    output_dir.mkdir(exist_ok=True)
    output_path = output_dir / f'{student.name}_小升初简历.docx'
    doc.save(output_path)
    
    return output_path

def export_to_word_green(student, base_dir):
    """导出为Word文档（绿色模板 - 基于简历模板1.pdf）"""
    doc = Document()
    
    # 绿色主题颜色
    GREEN = (91, 140, 62)  # #5B8C3E
    GREEN_LIGHT = (139, 195, 74)  # #8BC34A
    
    # 设置页面
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
    
    def add_green_title(text):
        """添加绿色标题"""
        para = doc.add_paragraph()
        para.paragraph_format.space_before = Pt(16)
        para.paragraph_format.space_after = Pt(10)
        run = para.add_run(f'  {text}  ')
        set_run_font(run, font_name='黑体', font_size=14, bold=True, color=(255, 255, 255))
        # 背景色需要通过shading设置
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), '5B8C3E')
        para._p.get_or_add_pPr().append(shd)
        return para
    
    # ========== 封面 ==========
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(120)
    run = para.add_run('我 的 简 历')
    set_run_font(run, font_name='华文行楷', font_size=42, bold=True, color=GREEN)
    
    # 照片
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(30)
    if student.photo:
        photo_path = base_dir / 'static' / student.photo
        if photo_path.exists():
            run = para.add_run()
            run.add_picture(str(photo_path), width=Cm(3.5))
    
    # 封面信息
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(20)
    run = para.add_run(student.name or '姓名')
    set_run_font(run, font_size=24, color=GREEN)
    
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run('─' * 20)
    set_run_font(run, font_size=10, color=GREEN)
    
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(f'◆ {student.school or "学校"}')
    set_run_font(run, font_size=12, color=(100, 100, 100))
    
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(f'◆ 电话：{student.phone or "联系电话"}')
    set_run_font(run, font_size=12, color=(100, 100, 100))
    
    doc.add_page_break()
    
    # ========== 个人简介 ==========
    add_green_title('个人简介')
    
    # 信息表格
    info_table = doc.add_table(rows=7, cols=3)
    info_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    
    info_items = [
        f'姓名：{student.name or ""}',
        f'性别：{student.gender or ""}',
        f'出生年月：{student.birth_date or ""}',
        f'毕业学校：{student.school or ""}',
        f'兴趣爱好：{student.hobbies or ""}',
        f'家庭住址：{student.address or student.district or ""}',
        f'联系电话：{student.phone or ""}',
    ]
    
    for i, info in enumerate(info_items):
        row = info_table.rows[i]
        cell = row.cells[0]
        cell.text = ''
        para = cell.paragraphs[0]
        run = para.add_run(info)
        set_run_font(run, font_size=11)
        
        # 照片在右侧（第一行合并）
        if i == 0:
            photo_cell = row.cells[2]
            for j in range(1, 7):
                photo_cell.merge(info_table.rows[j].cells[2])
            photo_para = photo_cell.paragraphs[0]
            photo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if student.photo:
                photo_path = base_dir / 'static' / student.photo
                if photo_path.exists():
                    run = photo_para.add_run()
                    run.add_picture(str(photo_path), width=Cm(3))
    
    # 成绩单
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(20)
    run = para.add_run('五年级成绩单')
    set_run_font(run, font_size=12, bold=True, color=GREEN)
    
    grade_table = doc.add_table(rows=3, cols=4)
    grade_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    grade_data = [
        ('科目', '语文', '英语', '数学'),
        ('第一学期', '优', '优', '优'),
        ('第二学期', '优', '优', '优'),
    ]
    for i, row_data in enumerate(grade_data):
        row = grade_table.rows[i]
        for j, text in enumerate(row_data):
            cell = row.cells[j]
            cell.text = ''
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(text)
            if i == 0:
                set_run_font(run, font_size=11, bold=True, color=(255, 255, 255))
                set_cell_background(cell, '5B8C3E')
            else:
                set_run_font(run, font_size=11)
    
    doc.add_page_break()
    
    # ========== 自荐书 ==========
    add_green_title('自荐书')
    
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(15)
    run = para.add_run('尊敬的XX学校老师：')
    set_run_font(run, font_size=11)
    
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Cm(0.75)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    intro_text = f'您好，我是{student.school or "XX小学"}的{student.name or "学生"}，我非常有意愿地就读贵校，希望您能使我梦想成真，非常感谢。下面说说我的自荐情况：'
    run = para.add_run(intro_text)
    set_run_font(run, font_size=11)
    
    if student.self_intro:
        for p_text in student.self_intro.split('\n'):
            if p_text.strip():
                para = doc.add_paragraph()
                para.paragraph_format.first_line_indent = Cm(0.75)
                para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
                run = para.add_run(p_text.strip())
                set_run_font(run, font_size=11)
    
    # 签名
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    para.paragraph_format.space_before = Pt(30)
    run = para.add_run('此致')
    set_run_font(run, font_size=11)
    
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = para.add_run('敬礼')
    set_run_font(run, font_size=11)
    
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    para.paragraph_format.space_before = Pt(20)
    run = para.add_run(f'自荐人：{student.school or ""} {student.class_name or ""} {student.name or ""}')
    set_run_font(run, font_size=11)
    
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = para.add_run(f'{datetime.now().year}年{datetime.now().month}月')
    set_run_font(run, font_size=11)
    
    doc.add_page_break()
    
    # ========== 获奖经历 ==========
    if student.awards:
        add_green_title('获奖经历')
        
        for award in sorted(student.awards, key=lambda x: x.order):
            para = doc.add_paragraph()
            para.paragraph_format.space_before = Pt(6)
            if award.date:
                run = para.add_run(f'{award.date}  ')
                set_run_font(run, font_size=11, bold=True, color=GREEN)
            run = para.add_run(award.title)
            set_run_font(run, font_size=11)
        
        doc.add_page_break()
    
    # ========== 证书 ==========
    certs = [c for c in student.certificates if c.category in ('证书', '奖状')]
    if certs:
        add_green_title('获奖证书')
        
        for i in range(0, len(certs), 2):
            table = doc.add_table(rows=2, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            for j in range(2):
                if i + j < len(certs):
                    cert = certs[i + j]
                    cell = table.cell(0, j)
                    para = cell.paragraphs[0]
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    
                    if cert.image_path:
                        img_path = base_dir / 'static' / cert.image_path
                        if img_path.exists():
                            run = para.add_run()
                            run.add_picture(str(img_path), width=Cm(6))
                    
                    cell = table.cell(1, j)
                    para = cell.paragraphs[0]
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = para.add_run(cert.name)
                    set_run_font(run, font_size=10, color=(100, 100, 100))
            
            doc.add_paragraph()
        
        doc.add_page_break()
    
    # ========== 寄语 ==========
    add_green_title('寄 语')
    
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(10)
    run = para.add_run('老师寄语 ' + '/' * 40)
    set_run_font(run, font_size=11, bold=True, color=GREEN)
    
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Cm(0.75)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    teacher_msg = student.teacher_comment or '你聪颖，你善良，你活泼。小小的你会长大，小小的你会成熟，愿你更坚强！愿你更自信！'
    run = para.add_run(teacher_msg)
    set_run_font(run, font_size=11)
    
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(20)
    run = para.add_run('家长寄语 ' + '/' * 40)
    set_run_font(run, font_size=11, bold=True, color=GREEN)
    
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Cm(0.75)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    parent_msg = student.parent_message or '孩子，相信自己，发挥自己的潜能！愿你生活与学习一帆风顺！一路阳光！'
    run = para.add_run(parent_msg)
    set_run_font(run, font_size=11)
    
    # 保存
    output_dir = base_dir / 'static' / 'exports'
    output_dir.mkdir(exist_ok=True)
    output_path = output_dir / f'{student.name}_小升初简历_绿色.docx'
    doc.save(output_path)
    
    return output_path


def export_to_word_template2(student, base_dir):
    """导出为Word文档（模板2 - 清新绿色版）"""
    doc = Document()
    
    # 绿色主题
    GREEN_MAIN = (67, 160, 71)  # #43A047
    GREEN_DARK = (46, 125, 50)  # #2E7D32
    
    # 设置页面
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
    
    def add_t2_title(text, icon='◆'):
        """添加绿色标题"""
        para = doc.add_paragraph()
        para.paragraph_format.space_before = Pt(20)
        para.paragraph_format.space_after = Pt(12)
        para.paragraph_format.border_bottom = True
        run = para.add_run(f'{icon} {text}')
        set_run_font(run, font_name='黑体', font_size=16, bold=True, color=GREEN_DARK)
        return para
    
    # ========== 封面 ==========
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(60)
    run = para.add_run('小升初简历')
    set_run_font(run, font_size=14, color=(100, 100, 100))
    
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(10)
    run = para.add_run('─' * 30)
    set_run_font(run, font_size=10, color=GREEN_MAIN)
    
    # 照片
    if student.photo:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(30)
        photo_path = base_dir / 'static' / student.photo
        if photo_path.exists():
            run = para.add_run()
            run.add_picture(str(photo_path), width=Cm(4))
    
    # 姓名
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(30)
    run = para.add_run(student.name or '姓名')
    set_run_font(run, font_name='黑体', font_size=36, bold=True, color=GREEN_DARK)
    
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(5)
    run = para.add_run('─' * 15)
    set_run_font(run, font_size=10, color=GREEN_MAIN)
    
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(15)
    run = para.add_run(student.school or '学校名称')
    set_run_font(run, font_size=14, color=(80, 80, 80))
    
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(8)
    run = para.add_run(f'联系电话：{student.phone or "电话"}')
    set_run_font(run, font_size=12, color=(100, 100, 100))
    
    doc.add_page_break()
    
    # ========== 基本信息 ==========
    add_t2_title('基本信息', '👤')
    
    info_table = doc.add_table(rows=4, cols=4)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    info_data = [
        ('姓名', student.name or '', '性别', student.gender or '男'),
        ('出生日期', student.birth_date or '', '身高', student.height or ''),
        ('毕业学校', student.school or '', '班级', student.class_name or ''),
        ('兴趣爱好', student.hobbies or '', '电话', student.phone or ''),
    ]
    
    for i, (l1, v1, l2, v2) in enumerate(info_data):
        row = info_table.rows[i]
        for j, (label, value) in enumerate([(l1, v1), (l2, v2)]):
            idx = j * 2
            # 标签
            cell = row.cells[idx]
            cell.text = ''
            para = cell.paragraphs[0]
            run = para.add_run(label)
            set_run_font(run, font_size=11, bold=True, color=GREEN_DARK)
            set_cell_background(cell, 'E8F5E9')
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            
            # 值
            cell = row.cells[idx + 1]
            cell.text = ''
            para = cell.paragraphs[0]
            run = para.add_run(value)
            set_run_font(run, font_size=11)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    
    # ========== 成绩单 ==========
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(20)
    run = para.add_run('学业成绩')
    set_run_font(run, font_size=13, bold=True, color=GREEN_DARK)
    
    grade_table = doc.add_table(rows=3, cols=4)
    grade_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    grade_data = [
        ('学期', '语文', '数学', '英语'),
        ('五年级上', '优', '优', '优'),
        ('五年级下', '优', '优', '优'),
    ]
    for i, row_data in enumerate(grade_data):
        row = grade_table.rows[i]
        for j, text in enumerate(row_data):
            cell = row.cells[j]
            cell.text = ''
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(text)
            if i == 0:
                set_run_font(run, font_size=11, bold=True, color=(255, 255, 255))
                set_cell_background(cell, '43A047')
            else:
                set_run_font(run, font_size=11)
                if i % 2 == 0:
                    set_cell_background(cell, 'E8F5E9')
    
    # ========== 自我介绍 ==========
    add_t2_title('自我介绍', '✏')
    
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Cm(0.75)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    intro_text = student.self_intro or '我是一个阳光、开朗、自信的学生，热爱学习，积极参加各种活动和比赛。在校期间成绩优异，多次获得各类荣誉。我相信通过自己的努力，一定能在中学阶段取得更好的成绩！'
    run = para.add_run(intro_text)
    set_run_font(run, font_size=11)
    
    # ========== 获奖经历 ==========
    if student.awards:
        add_t2_title('获奖经历', '🏆')
        
        for award in sorted(student.awards, key=lambda x: x.order):
            para = doc.add_paragraph()
            para.paragraph_format.space_before = Pt(8)
            para.paragraph_format.left_indent = Cm(0.5)
            
            # 日期标签
            if award.date:
                run = para.add_run(f'【{award.date}】')
                set_run_font(run, font_size=10, bold=True, color=GREEN_MAIN)
            
            run = para.add_run(f'  {award.title}')
            set_run_font(run, font_size=11)
    
    # ========== 证书照片 ==========
    if student.certificates:
        doc.add_page_break()
        add_t2_title('获奖证书', '📜')
        
        certs = list(student.certificates)
        for i in range(0, len(certs), 3):
            table = doc.add_table(rows=2, cols=3)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            for j in range(3):
                if i + j < len(certs):
                    cert = certs[i + j]
                    # 图片
                    cell = table.cell(0, j)
                    para = cell.paragraphs[0]
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    
                    if cert.image_path:
                        img_path = base_dir / 'static' / cert.image_path
                        if img_path.exists():
                            run = para.add_run()
                            run.add_picture(str(img_path), width=Cm(5))
                    
                    # 名称
                    cell = table.cell(1, j)
                    para = cell.paragraphs[0]
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = para.add_run(cert.name)
                    set_run_font(run, font_size=9, color=(100, 100, 100))
            
            doc.add_paragraph()
    
    # ========== 寄语 ==========
    doc.add_page_break()
    add_t2_title('寄语', '❤')
    
    if student.teacher_comment:
        para = doc.add_paragraph()
        para.paragraph_format.space_before = Pt(10)
        run = para.add_run('【老师寄语】')
        set_run_font(run, font_size=12, bold=True, color=GREEN_DARK)
        
        para = doc.add_paragraph()
        para.paragraph_format.first_line_indent = Cm(0.75)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        run = para.add_run(student.teacher_comment)
        set_run_font(run, font_size=11)
    
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(20)
    run = para.add_run('【家长寄语】')
    set_run_font(run, font_size=12, bold=True, color=GREEN_DARK)
    
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Cm(0.75)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    parent_msg = student.parent_message or '孩子，相信自己，发挥自己的潜能。愿你生活与学习一帆风顺！一路阳光！'
    run = para.add_run(parent_msg)
    set_run_font(run, font_size=11)
    
    # ========== 感谢 ==========
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(60)
    run = para.add_run('─' * 25)
    set_run_font(run, font_size=10, color=GREEN_MAIN)
    
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(20)
    run = para.add_run('感谢您的阅读')
    set_run_font(run, font_name='黑体', font_size=18, bold=True, color=GREEN_DARK)
    
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run('Thank You For Reading')
    set_run_font(run, font_size=11, color=(150, 150, 150))
    
    # 保存
    output_dir = base_dir / 'static' / 'exports'
    output_dir.mkdir(exist_ok=True)
    output_path = output_dir / f'{student.name}_小升初简历_模板2.docx'
    doc.save(output_path)
    
    return output_path


def export_to_pdf(student, base_dir):
    """导出为PDF（先生成Word再转换）"""
    import subprocess
    
    # 先生成Word
    word_path = export_to_word(student, base_dir)
    
    # 尝试使用libreoffice或word转换为PDF
    output_dir = base_dir / 'static' / 'exports'
    pdf_path = output_dir / f'{student.name}_小升初简历.pdf'
    
    try:
        # 尝试使用python-docx2pdf
        from docx2pdf import convert
        convert(str(word_path), str(pdf_path))
        return pdf_path
    except:
        pass
    
    try:
        # 尝试使用win32com
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(str(word_path.absolute()))
        doc.SaveAs(str(pdf_path.absolute()), FileFormat=17)  # 17 = PDF
        doc.Close()
        word.Quit()
        pythoncom.CoUninitialize()
        return pdf_path
    except:
        pass
    
    # 如果都失败，返回Word文件
    return word_path

