import os
import glob
import pandas as pd
from docx import Document
import sys

# Set encoding for stdout
sys.stdout.reconfigure(encoding='utf-8')

def get_profile_data(base_dir):
    profile_path = os.path.join(base_dir, 'profile.xlsx')
    data = {}
    if os.path.exists(profile_path):
        try:
            df = pd.read_excel(profile_path)
            # Flatten and simple extraction
            vals = df.values.flatten()
            keys = ['姓名', '性别', '出生日期', '身高', '学校', '电话', '住址']
            for k in keys:
                for i, val in enumerate(vals):
                    if str(val).replace(' ', '') == k and i + 1 < len(vals):
                        data[k] = str(vals[i+1]).strip()
        except Exception as e:
            print(f"Error reading profile: {e}")
    return data

def get_source_content(base_dir):
    pattern = os.path.join(base_dir, '*v3*修改版.docx')
    files = glob.glob(pattern)
    content = {
        'paragraphs': [],
        'tables': []
    }
    if files:
        try:
            doc = Document(files[0])
            for p in doc.paragraphs:
                if p.text.strip():
                    content['paragraphs'].append(p.text.strip())
            
            for table in doc.tables:
                t_data = []
                for row in table.rows:
                    r_data = [cell.text.strip() for cell in row.cells]
                    t_data.append(r_data)
                content['tables'].append(t_data)
        except Exception as e:
            print(f"Error reading source doc: {e}")
    return content

base_dir = os.path.join('study', 'resume')
profile = get_profile_data(base_dir)
print("--- PROFILE ---")
print(profile)

content = get_source_content(base_dir)
print("\n--- SOURCE CONTENT (First 5 paras) ---")
for p in content['paragraphs'][:5]:
    print(p)
print("\n--- SOURCE TABLES (First table) ---")
if content['tables']:
    for row in content['tables'][0]:
        print(row)

