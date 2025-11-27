import os
import glob
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import datetime

def set_font(run, font_name='微软雅黑', font_size=10.5, bold=False, color=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(font_size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def find_file(directory, pattern):
    files = glob.glob(os.path.join(directory, pattern))
    if files:
        return files[0]
    return None

def excel_date_to_str(serial):
    try:
        if isinstance(serial, (int, float)):
            dt = datetime.datetime(1899, 12, 30) + datetime.timedelta(days=serial)
            return dt.strftime('%Y年%m月')
        return str(serial)
    except:
        return str(serial)

def main():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    print(f"Working directory: {base_dir}")

    # 1. Load Profile
    profile_data = {}
    profile_path = os.path.join(base_dir, 'profile.xlsx')
    if os.path.exists(profile_path):
        try:
            df = pd.read_excel(profile_path)
            vals = df.values.flatten()
            # Basic mapping based on observed structure
            for i, val in enumerate(vals):
                s_val = str(val).strip()
                if '姓名' in s_val and i+1 < len(vals): profile_data['Name'] = str(vals[i+1]).strip()
                if '性别' in s_val and i+1 < len(vals): profile_data['Gender'] = str(vals[i+1]).strip()
                if '身高' in s_val and i+1 < len(vals): profile_data['Height'] = str(vals[i+1]).strip()
                if '出生' in s_val and i+1 < len(vals): profile_data['Birthday'] = excel_date_to_str(vals[i+1])
                if '学校' in s_val and i+1 < len(vals): profile_data['School'] = str(vals[i+1]).strip()
                if '电话' in s_val and i+1 < len(vals): profile_data['Phone'] = str(vals[i+1]).strip()
                if '住址' in s_val and i+1 < len(vals): profile_data['Address'] = str(vals[i+1]).strip()
        except Exception as e:
            print(f"Error reading profile: {e}")

    # 2. Load Source Resume Content
    source_path = find_file(base_dir, '*v3*修改版.docx')
    source_content = []
    if source_path:
        print(f"Reading source: {source_path}")
        try:
            src_doc = Document(source_path)
            for p in src_doc.paragraphs:
                if p.text.strip():
                    source_content.append(p.text.strip())
            # Also get tables from source if any
            for t in src_doc.tables:
                for r in t.rows:
                    row_txt = " ".join([c.text.strip() for c in r.cells if c.text.strip()])
                    if row_txt:
                        source_content.append(row_txt)
        except Exception as e:
            print(f"Error reading source doc: {e}")

    # 3. Load Template
    # Trying to match the user's request for "template 1"
    # Assuming (01)...docx is the docx version of template 1.doc
    template_path = find_file(os.path.join(base_dir, 'template'), '*01*word*.docx')
    if not template_path:
        # Fallback to any docx in template dir
        template_path = find_file(os.path.join(base_dir, 'template'), '*.docx')
    
    if not template_path:
        print("No suitable template found.")
        return

    print(f"Using template: {template_path}")
    doc = Document(template_path)

    # 4. Fill Template
    # We iterate through tables to find cells matching keys
    print("Filling template with profile data...")
    replacements = {
        '姓名': profile_data.get('Name', '赵Eric'),
        '性别': profile_data.get('Gender', '男'),
        '出生日期': profile_data.get('Birthday', '2013年12月'),
        '身高': profile_data.get('Height', '165cm'),
        '所在小学': profile_data.get('School', '（请填写学校）'),
        '联系电话': profile_data.get('Phone', '（请填写）'),
        '家庭住址': profile_data.get('Address', '（请填写）'),
        'Name': profile_data.get('Name', '赵Eric'), # English key fallback
    }

    for table in doc.tables:
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                txt = cell.text.strip()
                # Exact match or contains
                for key, val in replacements.items():
                    if key in txt and len(txt) < 10: # Simple label cell
                        # Try to find the value cell (next one usually)
                        if i + 1 < len(row.cells):
                            # Don't overwrite if it already has data, unless it's a placeholder
                            val_cell = row.cells[i+1]
                            if not val_cell.text.strip() or '填写' in val_cell.text:
                                val_cell.text = val
                                # Reset font for the new text
                                for p in val_cell.paragraphs:
                                    for r in p.runs:
                                        set_font(r, font_size=11)
    
    # 5. Append Source Content
    # We append the detailed text from the v3 resume at the end of the document
    # or in a "Personal Description" section if found.
    
    # Check if there is a "Self Introduction" section in template to fill?
    # For now, we append cleanly at the end or replace a placeholder text block.
    
    doc.add_page_break()
    heading = doc.add_paragraph()
    run = heading.add_run("详细个人介绍与经历")
    set_font(run, font_size=14, bold=True, color=(0, 112, 192))
    
    for line in source_content:
        # Filter out basic info lines we already used
        if any(k in line for k in ['姓名', '性别', '身高']):
            continue
            
        p = doc.add_paragraph()
        run = p.add_run(line)
        set_font(run, font_size=10.5)
        # Indent slightly
        p.paragraph_format.left_indent = Cm(0.5)

    # 6. Add Photos
    cert_dir = os.path.join(base_dir, 'certificate')
    photos = []
    if os.path.exists(cert_dir):
        # Look in root and bak
        for root, dirs, files in os.walk(cert_dir):
            for f in files:
                if f.lower().endswith(('.jpg', '.png', '.jpeg')):
                    photos.append(os.path.join(root, f))
    
    if photos:
        doc.add_page_break()
        heading = doc.add_paragraph()
        run = heading.add_run("证书与风采展示")
        set_font(run, font_size=14, bold=True, color=(0, 112, 192))
        
        # Grid layout for photos: 2 per row
        table = doc.add_table(rows=0, cols=2)
        table.autofit = True
        
        row_cells = None
        for i, photo_path in enumerate(photos):
            if i % 2 == 0:
                row_cells = table.add_row().cells
            
            cell = row_cells[i % 2]
            p = cell.paragraphs[0]
            run = p.add_run()
            try:
                run.add_picture(photo_path, width=Inches(2.5))
                # Add caption
                filename = os.path.basename(photo_path)
                p_cap = cell.add_paragraph(filename)
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_font(p_cap.runs[0] if p_cap.runs else p_cap.add_run(filename), font_size=9)
            except Exception as e:
                print(f"Could not add photo {photo_path}: {e}")

    # 7. Save
    output_file = os.path.join(base_dir, '赵Eric_小升初简历_新模板_2.docx')
    doc.save(output_file)
    print(f"Generated: {output_file}")

if __name__ == '__main__':
    main()

