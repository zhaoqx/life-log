import os
import glob
from docx import Document
import sys

# Set encoding for stdout
sys.stdout.reconfigure(encoding='utf-8')

def analyze_docx_stream(path):
    print(f"Opening: {path}")
    try:
        with open(path, 'rb') as f:
            doc = Document(f)
            print("--- PARAGRAPHS ---")
            count = 0
            for p in doc.paragraphs:
                if p.text.strip():
                    print(f"P: {p.text.strip()[:50]}")
                    count += 1
            print(f"Total paragraphs: {count}")
            
            print("\n--- TABLES ---")
            for i, table in enumerate(doc.tables):
                print(f"Table {i}: {len(table.rows)} rows")
                for r in table.rows:
                    txt = " ".join([c.text.strip() for c in r.cells])
                    if txt.strip():
                        print(f"  {txt[:50]}")
    except Exception as e:
        print(f"Error: {e}")

resume_dir = os.path.join('study', 'resume')
# Find template
template_pattern = os.path.join(resume_dir, 'template', '*01*word*.docx')
templates = glob.glob(template_pattern)
if templates:
    analyze_docx_stream(templates[0])
else:
    print("Template not found")

# Find source
source_pattern = os.path.join(resume_dir, '*v3*修改版.docx')
sources = glob.glob(source_pattern)
if sources:
    analyze_docx_stream(sources[0])
else:
    print("Source not found")

