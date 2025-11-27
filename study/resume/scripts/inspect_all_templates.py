import os
import glob
from docx import Document
import sys

# Set encoding for stdout
sys.stdout.reconfigure(encoding='utf-8')

def analyze_template(path):
    print(f"Analyzing template: {path}")
    try:
        doc = Document(path)
        print(f"  Paragraphs: {len(doc.paragraphs)}")
        print(f"  Tables: {len(doc.tables)}")
        if doc.tables:
            print(f"  First table rows: {len(doc.tables[0].rows)}")
            if len(doc.tables[0].rows) > 0:
                print(f"  First row text: {[c.text for c in doc.tables[0].rows[0].cells]}")
    except Exception as e:
        print(f"  Error: {e}")

base_dir = os.path.join('study', 'resume', 'template')
files = glob.glob(os.path.join(base_dir, '*.docx'))

for f in files:
    analyze_template(f)

