import os
import glob
from docx import Document
import sys

# Set encoding for stdout
sys.stdout.reconfigure(encoding='utf-8')

def find_file(directory, pattern):
    files = glob.glob(os.path.join(directory, pattern))
    if files:
        return files[0]
    return None

def analyze_docx(path):
    print(f"Analyzing: {path}")
    if not path or not os.path.exists(path):
        print("File not found.")
        return

    try:
        doc = Document(path)
        print("--- PARAGRAPHS ---")
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip():
                print(f"P{i}: {p.text.strip()[:100]}")
        
        print("\n--- TABLES ---")
        for i, table in enumerate(doc.tables):
            print(f"Table {i}: {len(table.rows)} rows x {len(table.columns)} cols")
            for r_idx, row in enumerate(table.rows):
                row_text = " | ".join([cell.text.strip() for cell in row.cells])
                if row_text.strip():
                    print(f"  R{r_idx}: {row_text[:100]}")
    except Exception as e:
        print(f"Error: {e}")

resume_dir = os.path.join('study', 'resume')
template_path = os.path.join(resume_dir, 'template', '（01）升学简历word【小升初】.docx')
source_path = find_file(resume_dir, '*v3*修改版.docx')
print(f"Found source: {source_path}")

analyze_docx(template_path)
print("\n" + "="*50 + "\n")
if source_path:
    analyze_docx(source_path)
