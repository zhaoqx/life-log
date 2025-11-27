import os
import glob
import json
import pandas as pd
from PIL import Image
import datetime
import shutil
import sys

# --- Configuration ---
BASE_DIR = os.path.join('study', 'resume')
WEB_BUILD_DIR = os.path.join(BASE_DIR, 'web_build')
ASSETS_DIR = os.path.join(WEB_BUILD_DIR, 'assets')
DATA_FILE = os.path.join(WEB_BUILD_DIR, 'data.json')

if not os.path.exists(ASSETS_DIR):
    os.makedirs(ASSETS_DIR)

# --- Data Extraction ---
def load_and_save_data():
    data = {
        'meta': {
            'generated_at': datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'template': '小升初简历模板1'
        },
        'profile': {},
        'sections': {
            'specialty': [],
            'awards': [],
            'activities': [],
            'intro': [],
            'parents_words': []
        },
        'grades': [
            {'term': '五年级上学期', 'chinese': '100', 'math': '96', 'english': '98'},
            {'term': '五年级下学期', 'chinese': '98', 'math': '98', 'english': '96'}
        ],
        'attachments': []
    }
    
    # 1. Profile from Excel
    profile_path = os.path.join(BASE_DIR, 'profile.xlsx')
    if os.path.exists(profile_path):
        try:
            df = pd.read_excel(profile_path)
            vals = df.values.flatten()
            for i, val in enumerate(vals):
                s = str(val).strip().replace(' ', '')
                if s == '姓名' and i+1 < len(vals): data['profile']['name'] = str(vals[i+1])
                if s == '性别' and i+1 < len(vals): data['profile']['gender'] = str(vals[i+1])
                if s == '出生日期' and i+1 < len(vals): 
                    d = vals[i+1]
                    if isinstance(d, (int, float)):
                        dt = datetime.datetime(1899, 12, 30) + datetime.timedelta(days=d)
                        data['profile']['birthday'] = dt.strftime('%Y年%m月')
                    else:
                        data['profile']['birthday'] = str(d)
                if s == '身高' and i+1 < len(vals): data['profile']['height'] = str(vals[i+1])
                # Add placeholders if missing
                data['profile']['school'] = data['profile'].get('school', '（请填写学校）')
                data['profile']['phone'] = data['profile'].get('phone', '（请填写电话）')
                data['profile']['address'] = data['profile'].get('address', '（请填写地址）')
                data['profile']['hometown'] = data['profile'].get('hometown', '（请填写户籍）')
        except Exception as e:
            print(f"Error reading profile: {e}")

    # 2. Content from Word
    from docx import Document
    v3_files = glob.glob(os.path.join(BASE_DIR, '*v3*修改版.docx'))
    raw_lines = []
    if v3_files:
        try:
            doc = Document(v3_files[0])
            for p in doc.paragraphs:
                if p.text.strip():
                    raw_lines.append(p.text.strip())
        except:
            pass
            
    # Classify Content
    for line in raw_lines:
        if len(line) < 2: continue
        if any(k in line for k in ['姓名', '性别', '表格', '出生']): continue
        
        if '围棋' in line or '特长' in line:
            data['sections']['specialty'].append(line)
        elif '奖' in line or '荣誉' in line or '证书' in line or '杯' in line:
            data['sections']['awards'].append(line)
        elif '足球' in line or '活动' in line or '志愿者' in line:
            data['sections']['activities'].append(line)
        elif len(line) > 30: 
            data['sections']['intro'].append(line)
            
    # Default Parents Words
    data['sections']['parents_words'].append('（此处可由家长填写对孩子的评价和期望）')

    # 3. Process Images
    cert_dir = os.path.join(BASE_DIR, 'certificate')
    if os.path.exists(cert_dir):
        for root, _, files in os.walk(cert_dir):
            for f in files:
                if f.lower().endswith(('.jpg', '.jpeg', '.png')):
                    src_path = os.path.join(root, f)
                    safe_name = f.replace(' ', '_')
                    dst_rel_path = f'assets/{safe_name}'
                    dst_path = os.path.join(WEB_BUILD_DIR, dst_rel_path)
                    
                    try:
                        with Image.open(src_path) as img:
                            # Compress
                            max_width = 800
                            if img.width > max_width:
                                ratio = max_width / img.width
                                new_height = int(img.height * ratio)
                                img = img.resize((max_width, new_height), Image.Resampling.LANCZOS)
                            if img.mode in ("RGBA", "P"):
                                img = img.convert("RGB")
                            img.save(dst_path, quality=70, optimize=True)
                            
                            data['attachments'].append({
                                'name': os.path.splitext(f)[0],
                                'path': dst_rel_path,
                                'type': 'image'
                            })
                    except Exception as e:
                        print(f"Error processing image {f}: {e}")

    # Save Data
    with open(DATA_FILE, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        
    return data

# --- HTML Generation ---
def generate_html(data):
    css = """
    <style>
        body {
            background: #525659;
            font-family: "Microsoft YaHei", "SimHei", sans-serif;
            margin: 0;
            padding: 20px;
            display: flex;
            flex-direction: column;
            align-items: center;
        }
        .page {
            background: white;
            width: 210mm;
            min-height: 297mm;
            padding: 2.54cm 2cm; /* Word default margins approx */
            box-shadow: 0 0 10px rgba(0,0,0,0.5);
            box-sizing: border-box;
            margin-bottom: 20px;
            position: relative;
        }
        @media print {
            body { background: none; margin: 0; padding: 0; display: block; }
            .page { width: 100%; box-shadow: none; margin: 0; padding: 2.54cm 2cm; page-break-after: always; height: auto; min-height: 0; }
            .no-print { display: none !important; }
        }
        
        /* Resume Header */
        .resume-title {
            text-align: center;
            color: #0070c0; /* Blue standard */
            font-size: 26pt;
            font-weight: bold;
            margin: 0 0 10px 0;
            letter-spacing: 5px;
        }
        .header-line {
            border-top: 2px solid #ccc;
            margin-bottom: 20px;
            width: 100%;
        }
        
        /* Section Titles */
        .section-title {
            color: #0070c0;
            font-size: 14pt;
            font-weight: bold;
            margin-top: 15px;
            margin-bottom: 10px;
            display: flex;
            align-items: center;
        }
        .section-title::before {
            content: "◆";
            margin-right: 5px;
            font-size: 12pt;
        }
        
        /* Tables */
        table {
            width: 100%;
            border-collapse: collapse;
            margin-bottom: 15px;
            border: 1px solid #000;
        }
        td, th {
            border: 1px solid #000;
            padding: 8px 5px;
            font-size: 11pt;
            vertical-align: middle;
        }
        
        /* Info Table Specifics */
        .info-label {
            background-color: #e8f4fc; /* Light blue */
            color: #0070c0;
            font-weight: bold;
            text-align: center;
            width: 15%;
        }
        .info-value {
            text-align: center;
            width: 35%;
        }
        
        /* Grade Table Specifics */
        .grade-header {
            background-color: #0070c0;
            color: white;
            font-weight: bold;
            text-align: center;
        }
        .grade-term {
            background-color: #e8f4fc;
            text-align: center;
        }
        .grade-val {
            text-align: center;
        }
        
        /* Lists */
        ul {
            list-style: none;
            padding-left: 0;
            margin: 5px 0;
        }
        li {
            position: relative;
            padding-left: 20px;
            margin-bottom: 5px;
            font-size: 11pt;
            line-height: 1.5;
        }
        li::before {
            content: "◇";
            position: absolute;
            left: 0;
            color: #000; /* Usually black bullet in this template, or blue? Let's stick to black/dark grey */
        }
        
        /* Text Blocks */
        .text-block {
            font-size: 11pt;
            line-height: 1.6;
            text-indent: 2em;
            text-align: justify;
        }
        
        .parent-block {
            color: #888;
            font-style: italic;
        }
        
        /* Photo Grid */
        .photo-grid {
            display: grid;
            grid-template-columns: 1fr 1fr;
            gap: 20px;
        }
        .photo-card {
            border: 1px solid #ddd;
            padding: 10px;
            text-align: center;
        }
        .photo-card img {
            max-width: 100%;
            height: 200px;
            object-fit: contain;
        }
        .photo-name {
            margin-top: 5px;
            font-size: 10pt;
            color: #555;
        }
        
        /* Floating Controls */
        .toolbar {
            position: fixed;
            top: 20px;
            right: 20px;
            background: white;
            padding: 10px;
            border-radius: 5px;
            box-shadow: 0 2px 10px rgba(0,0,0,0.2);
            z-index: 100;
        }
        .btn {
            background: #0070c0;
            color: white;
            border: none;
            padding: 8px 15px;
            border-radius: 3px;
            cursor: pointer;
            margin-left: 5px;
        }
        .btn:hover { background: #005294; }
    </style>
    """
    
    # Build Sections HTML
    specialty_html = "".join(f"<li>{item}</li>" for item in data['sections']['specialty'])
    if not specialty_html: specialty_html = "<li>围棋业余二段</li>"
    
    awards_html = "".join(f"<li>{item}</li>" for item in data['sections']['awards'])
    if not awards_html: awards_html = "<li>（获奖经历）</li>"
    
    activities_html = "".join(f"<li>{item}</li>" for item in data['sections']['activities'])
    if not activities_html: activities_html = "<li>（活动经历）</li>"
    
    intro_text = "<br>".join(data['sections']['intro']) if data['sections']['intro'] else "（自我介绍内容）"
    parents_text = data['sections']['parents_words'][0] if data['sections']['parents_words'] else ""

    # Build Grade Rows
    grade_rows = ""
    for g in data['grades']:
        grade_rows += f"""
        <tr>
            <td class="grade-term">{g['term']}</td>
            <td class="grade-val">{g['chinese']}</td>
            <td class="grade-val">{g['math']}</td>
            <td class="grade-val">{g['english']}</td>
        </tr>
        """

    # Build Attachments HTML
    attachments_html = ""
    for att in data['attachments']:
        attachments_html += f"""
        <div class="photo-card">
            <img src="{att['path']}" alt="{att['name']}">
            <div class="photo-name">{att['name']}</div>
        </div>
        """

    html = f"""
    <!DOCTYPE html>
    <html lang="zh-CN">
    <head>
        <meta charset="UTF-8">
        <meta name="viewport" content="width=device-width, initial-scale=1.0">
        <title>{data['profile'].get('name', '简历')} - 小升初</title>
        {css}
    </head>
    <body>
        <div class="toolbar no-print">
            <button class="btn" onclick="window.print()">打印 / 下载PDF</button>
            <a href="data.json" target="_blank" class="btn" style="text-decoration:none;">查看源数据</a>
        </div>

        <!-- Page 1: Main Resume -->
        <div class="page">
            <div class="resume-title">【 个 人 简 历 】</div>
            <div class="header-line"></div>
            
            <div class="section-title">基本信息</div>
            <table>
                <tr>
                    <td class="info-label">姓 名</td>
                    <td class="info-value">{data['profile'].get('name', '')}</td>
                    <td class="info-label">性 别</td>
                    <td class="info-value">{data['profile'].get('gender', '')}</td>
                </tr>
                <tr>
                    <td class="info-label">出生日期</td>
                    <td class="info-value">{data['profile'].get('birthday', '')}</td>
                    <td class="info-label">身 高</td>
                    <td class="info-value">{data['profile'].get('height', '')}</td>
                </tr>
                <tr>
                    <td class="info-label">所在小学</td>
                    <td class="info-value">{data['profile'].get('school', '')}</td>
                    <td class="info-label">户籍所在</td>
                    <td class="info-value">{data['profile'].get('hometown', '')}</td>
                </tr>
                <tr>
                    <td class="info-label">联系电话</td>
                    <td class="info-value">{data['profile'].get('phone', '')}</td>
                    <td class="info-label">家庭住址</td>
                    <td class="info-value">{data['profile'].get('address', '')}</td>
                </tr>
            </table>
            
            <div class="section-title">学业成绩</div>
            <table>
                <tr class="grade-header">
                    <td>学期</td>
                    <td>语文</td>
                    <td>数学</td>
                    <td>英语</td>
                </tr>
                {grade_rows}
            </table>
            
            <div class="section-title">特长与证书</div>
            <ul>{specialty_html}</ul>
            
            <div class="section-title">荣誉奖励</div>
            <ul>{awards_html}</ul>
            
            <div class="section-title">课外活动</div>
            <ul>{activities_html}</ul>
            
            <div class="section-title">兴趣爱好</div>
            <div class="text-block">围棋、足球、阅读、写作</div>
            
            <div class="section-title">自我介绍</div>
            <div class="text-block">{intro_text}</div>
            
            <div class="section-title">家长寄语</div>
            <div class="text-block parent-block">{parents_text}</div>
        </div>
        
        <!-- Page 2: Attachments -->
        <div class="page">
             <div class="resume-title" style="font-size:20pt;">【 证 书 与 风 采 】</div>
             <div class="header-line"></div>
             <div class="photo-grid">
                {attachments_html}
             </div>
        </div>
    </body>
    </html>
    """
    
    out_path = os.path.join(WEB_BUILD_DIR, 'index.html')
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write(html)
    return out_path

def main():
    print("Initializing Web Build...")
    data = load_and_save_data()
    print("Generating HTML...")
    path = generate_html(data)
    print(f"Success! Web site ready at: {path}")

if __name__ == '__main__':
    main()

