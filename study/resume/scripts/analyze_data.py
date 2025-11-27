import os
import glob
import pandas as pd
from docx import Document
import sys

sys.stdout.reconfigure(encoding='utf-8')

base_dir = os.path.join('study', 'resume')

def read_profile():
    path = os.path.join(base_dir, 'profile.xlsx')
    if not os.path.exists(path):
        print("Profile not found")
        return
    print(f"--- PROFILE ({path}) ---")
    try:
        df = pd.read_excel(path)
        print(df.to_string())
    except Exception as e:
        print(f"Error reading profile: {e}")

def read_docx(pattern, name):
    files = glob.glob(os.path.join(base_dir, pattern))
    if not files:
        files = glob.glob(os.path.join(base_dir, 'template', pattern))
    
    if files:
        path = files[0]
        print(f"\n--- {name} ({path}) ---")
        try:
            doc = Document(path)
            print("Paragraphs:")
            for i, p in enumerate(doc.paragraphs):
                if p.text.strip():
                    print(f"  P{i}: {p.text.strip()[:50]}")
            print("Tables:")
            for i, table in enumerate(doc.tables):
                print(f"  Table {i}:")
                for r_idx, row in enumerate(table.rows):
                    row_txt = " | ".join([c.text.strip() for c in row.cells])
                    if row_txt.strip():
                        print(f"    R{r_idx}: {row_txt[:80]}")
        except Exception as e:
            print(f"Error reading docx: {e}")
    else:
        print(f"\n{name} not found with pattern: {pattern}")

read_profile()
read_docx('*v3*修改版.docx', 'SOURCE RESUME')
read_docx('*01*升学*.docx', 'TEMPLATE')

