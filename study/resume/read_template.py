from docx import Document
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

template_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'template')

# 读取所有模板
templates = os.listdir(template_dir)

for template_name in templates:
    if template_name.endswith('.docx'):
        template_file = os.path.join(template_dir, template_name)
        print(f"\n{'='*50}")
        print(f"=== 读取模板: {template_name} ===")
        print('='*50)
        
        try:
            doc = Document(template_file)
            
            # 读取段落
            print("\n--- 段落内容 ---")
            para_count = 0
            for i, para in enumerate(doc.paragraphs):
                if para.text.strip():
                    print(f"段落{i}: {para.text[:100]}...")
                    para_count += 1
                    if para_count > 10:
                        print("... (更多段落省略)")
                        break

            # 读取表格
            print(f"\n--- 表格数量: {len(doc.tables)} ---")
            for t_idx, table in enumerate(doc.tables):
                print(f"\n表格 {t_idx} (行数: {len(table.rows)})")
                for r_idx, row in enumerate(table.rows):
                    if r_idx > 5:
                        print("... (更多行省略)")
                        break
                    row_text = []
                    for cell in row.cells:
                        text = cell.text.strip().replace('\n', ' ')[:30]
                        row_text.append(text)
                    print(f"  行{r_idx}: {' | '.join(row_text)}")
        except Exception as e:
            print(f"错误: {e}")

