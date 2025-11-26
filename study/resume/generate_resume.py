# -*- coding: utf-8 -*-
"""
小升初简历生成脚本
"""
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

def set_cell_background(cell, color):
    """设置单元格背景色"""
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_run_font(run, font_name='微软雅黑', font_size=11, bold=False, color=None):
    """设置字体"""
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(font_size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(color[0], color[1], color[2])

def add_heading_paragraph(doc, text, font_size=22, bold=True, color=(0, 112, 192), align='center'):
    """添加标题段落"""
    para = doc.add_paragraph()
    if align == 'center':
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'left':
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run(text)
    set_run_font(run, font_size=font_size, bold=bold, color=color)
    return para

def add_section_title(doc, text):
    """添加章节标题"""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run('◆ ' + text)
    set_run_font(run, font_size=14, bold=True, color=(0, 112, 192))
    para.paragraph_format.space_before = Pt(12)
    para.paragraph_format.space_after = Pt(6)
    return para

def create_resume():
    """创建小升初简历"""
    doc = Document()
    
    # 设置页面边距
    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
    
    # =============== 简历标题 ===============
    add_heading_paragraph(doc, '【 个 人 简 历 】', font_size=28, bold=True, color=(0, 82, 148))
    
    # 添加分隔线
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run('━' * 40)
    set_run_font(run, font_size=10, color=(200, 200, 200))
    
    # =============== 基本信息 ===============
    add_section_title(doc, '基本信息')
    
    # 创建基本信息表格
    table = doc.add_table(rows=4, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # 表格数据
    info_data = [
        ('姓    名', '赵Eric', '性    别', '男'),
        ('出生日期', '2013年12月', '身    高', '165cm'),
        ('所在小学', '（请填写学校名称）', '户籍所在', '（请填写）'),
        ('联系电话', '（请填写）', '家庭住址', '（请填写）'),
    ]
    
    for i, row_data in enumerate(info_data):
        row = table.rows[i]
        for j, text in enumerate(row_data):
            cell = row.cells[j]
            cell.text = ''
            para = cell.paragraphs[0]
            run = para.add_run(text)
            if j % 2 == 0:  # 标签列
                set_run_font(run, font_size=11, bold=True, color=(0, 82, 148))
                set_cell_background(cell, 'E8F4FC')
            else:  # 值列
                set_run_font(run, font_size=11, bold=False)
    
    # =============== 学业成绩 ===============
    add_section_title(doc, '学业成绩')
    
    # 创建成绩表格
    grade_table = doc.add_table(rows=3, cols=4)
    grade_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    grade_data = [
        ('学期', '语文', '数学', '英语'),
        ('五年级上学期', '100', '96', '98'),
        ('五年级下学期', '98', '98', '96'),
    ]
    
    for i, row_data in enumerate(grade_data):
        row = grade_table.rows[i]
        for j, text in enumerate(row_data):
            cell = row.cells[j]
            cell.text = ''
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(text)
            if i == 0:  # 表头
                set_run_font(run, font_size=11, bold=True, color=(255, 255, 255))
                set_cell_background(cell, '0070C0')
            else:
                set_run_font(run, font_size=11, bold=False)
                if j == 0:
                    set_cell_background(cell, 'E8F4FC')
    
    # =============== 特长与证书 ===============
    add_section_title(doc, '特长与证书')
    
    para = doc.add_paragraph()
    run = para.add_run('◇ 围棋业余二段')
    set_run_font(run, font_size=11)
    para.paragraph_format.left_indent = Cm(0.5)
    
    para = doc.add_paragraph()
    run = para.add_run('  围棋培养了逻辑思维能力和耐心，在对弈中学会了冷静分析和长远规划。')
    set_run_font(run, font_size=10, color=(100, 100, 100))
    para.paragraph_format.left_indent = Cm(0.5)
    
    # =============== 荣誉奖励 ===============
    add_section_title(doc, '荣誉奖励')
    
    para = doc.add_paragraph()
    run = para.add_run('◇ 钟书阁读书会获奖作文')
    set_run_font(run, font_size=11)
    para.paragraph_format.left_indent = Cm(0.5)
    
    para = doc.add_paragraph()
    run = para.add_run('  热爱阅读和写作，积极参加读书活动，作文获得钟书阁读书会奖项认可。')
    set_run_font(run, font_size=10, color=(100, 100, 100))
    para.paragraph_format.left_indent = Cm(0.5)
    
    # =============== 课外活动 ===============
    add_section_title(doc, '课外活动')
    
    para = doc.add_paragraph()
    run = para.add_run('◇ 校足球队队员')
    set_run_font(run, font_size=11)
    para.paragraph_format.left_indent = Cm(0.5)
    
    para = doc.add_paragraph()
    run = para.add_run('  热爱体育运动，作为校足球队队员，培养了团队合作精神和拼搏意识，增强了身体素质。')
    set_run_font(run, font_size=10, color=(100, 100, 100))
    para.paragraph_format.left_indent = Cm(0.5)
    
    # =============== 兴趣爱好 ===============
    add_section_title(doc, '兴趣爱好')
    
    para = doc.add_paragraph()
    run = para.add_run('围棋、足球、阅读、写作')
    set_run_font(run, font_size=11)
    para.paragraph_format.left_indent = Cm(0.5)
    
    # =============== 自我介绍 ===============
    add_section_title(doc, '自我介绍')
    
    intro_text = """    我是一名阳光开朗的小学生，品学兼优，全面发展。

    在学习方面，我认真努力，成绩优异，语文、数学、英语三科均保持在96分以上。我热爱阅读，曾在钟书阁读书会中获得作文奖项，阅读和写作是我最大的爱好之一。

    在特长方面，我学习围棋已达到业余二段水平，围棋不仅锻炼了我的逻辑思维能力，也让我学会了在复杂局面中保持冷静和耐心。

    在体育方面，我是校足球队的一员，足球运动让我懂得了团队合作的重要性，也培养了我不怕困难、勇于拼搏的精神。

    我希望能够进入贵校学习，在更好的平台上继续努力，全面提升自己，成为德智体美劳全面发展的好学生！"""
    
    para = doc.add_paragraph()
    run = para.add_run(intro_text)
    set_run_font(run, font_size=11)
    para.paragraph_format.line_spacing = 1.5
    
    # =============== 家长寄语 ===============
    add_section_title(doc, '家长寄语')
    
    parent_text = """    （此处可由家长填写对孩子的评价和期望）"""
    
    para = doc.add_paragraph()
    run = para.add_run(parent_text)
    set_run_font(run, font_size=11, color=(150, 150, 150))
    
    # 保存文档
    output_dir = os.path.dirname(os.path.abspath(__file__))
    output_file = os.path.join(output_dir, '赵Eric_小升初简历.docx')
    doc.save(output_file)
    print(f'简历已生成: {output_file}')
    return output_file

if __name__ == '__main__':
    create_resume()

